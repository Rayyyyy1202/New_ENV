#!/usr/bin/env python3
# 生成销售用 PPT(.pptx)· Amazon 对标选品系统
# 运行:python3 scripts/make-ppt.py
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK = RGBColor(0x16, 0x18, 0x1D)
MUTED = RGBColor(0x4B, 0x52, 0x60)
FAINT = RGBColor(0x9A, 0xA0, 0xAC)
ACCENT = RGBColor(0x5B, 0x5B, 0xD6)
ACCENTSOFT = RGBColor(0xEE, 0xF0, 0xFB)
OK = RGBColor(0x16, 0xA3, 0x4A)
DANGER = RGBColor(0xDC, 0x26, 0x26)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CANVAS = RGBColor(0xF6, 0xF7, 0xF9)
LINE = RGBColor(0xE3, 0xE6, 0xEC)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def _set_ea(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, l, t, w, h, fill=None, line=None, radius=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, l, t, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def text(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    # runs: list of paragraphs; each paragraph = list of (txt, size, color, bold)
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        for (txt, size, color, bold) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = FONT
            _set_ea(r)
    return tb


def footer(s, n):
    text(s, Inches(0.6), Inches(7.0), Inches(8), Inches(0.4),
         [[("LinkFinder · AI 对标选品系统", 9, FAINT, False)]])
    text(s, Inches(11.5), Inches(7.0), Inches(1.3), Inches(0.4),
         [[(f"{n:02d}", 9, FAINT, False)]], align=PP_ALIGN.RIGHT)


def header(s, kicker, title):
    rect(s, Inches(0.6), Inches(0.6), Inches(0.12), Inches(0.5), fill=ACCENT)
    text(s, Inches(0.85), Inches(0.55), Inches(11), Inches(0.4),
         [[(kicker, 12, ACCENT, True)]])
    text(s, Inches(0.83), Inches(0.95), Inches(11.6), Inches(0.9),
         [[(title, 30, INK, True)]])


def bullets(s, items, l=Inches(0.9), t=Inches(2.1), w=Inches(11.5), size=16, gap=10):
    runs = []
    for it in items:
        if isinstance(it, tuple):
            head, sub = it
            runs.append([("●  ", 14, ACCENT, True), (head, size, INK, True)])
            if sub:
                runs.append([("     " + sub, size - 3, MUTED, False)])
        else:
            runs.append([("●  ", 14, ACCENT, True), (it, size, INK, False)])
    text(s, l, t, w, Inches(4.5), runs, space=gap)


def card(s, l, t, w, h, title, body, num=None):
    c = rect(s, l, t, w, h, fill=WHITE, line=LINE, radius=True)
    badge = ("  " + num + "   ") if num else ""
    runs = [[(badge, 13, ACCENT, True), (title, 16, INK, True)], [(body, 12.5, MUTED, False)]]
    text(s, l + Inches(0.25), t + Inches(0.22), w - Inches(0.5), h - Inches(0.4), runs, space=8)
    return c


def chip(s, l, t, w, txt, fill=ACCENTSOFT, fg=ACCENT):
    h = Inches(0.62)
    rect(s, l, t, w, h, fill=fill, radius=True)
    text(s, l, t, w, h, [[(txt, 12, fg, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return h


# ───────── S1 封面 ─────────
s = slide(ACCENT)
rect(s, 0, 0, SW, SH, fill=ACCENT)
text(s, Inches(1.0), Inches(2.3), Inches(11), Inches(0.5),
     [[("AI 驱动 · 不用 RPA · Keepa 数据 + 大模型", 15, RGBColor(0xD6, 0xD9, 0xF7), True)]])
text(s, Inches(0.95), Inches(2.8), Inches(11.4), Inches(1.6),
     [[("Amazon 对标选品系统", 46, WHITE, True)]])
text(s, Inches(1.0), Inches(4.3), Inches(11), Inches(1.0),
     [[("上传报关单,AI 自动推荐可用的亚马逊对标链接,并讲清每一条理由 —— 越用越准", 18, RGBColor(0xE7, 0xE9, 0xFb), False)]])
text(s, Inches(1.0), Inches(6.4), Inches(11), Inches(0.5),
     [[("销售方案 · [公司名称]", 13, RGBColor(0xC7, 0xCB, 0xF2), False)]])

# ───────── S2 痛点 ─────────
s = slide()
header(s, "现状 · 痛点", "人工对标,有多难")
bullets(s, [
    "看懂中文品名 → 翻成英文 → 在历史库里翻找有没有现成对标",
    "没有就上亚马逊德国,用 4–5 种排序一个个搜、一张张点开看",
    "每个候选还要去欧盟商标库查品牌、核对材质、核对原产地(必须中国)",
    "一个商品 10–20 分钟:繁琐、跨多网站、大量重复判断、易漏易错、合规有风险",
])
rect(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.9), fill=ACCENTSOFT, radius=True)
text(s, Inches(1.2), Inches(5.6), Inches(11), Inches(0.9),
     [[("这正是 AI + 数据最该替代的环节:判断多、重复高、跨系统。", 16, ACCENT, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# ───────── S3 方案价值 ─────────
s = slide()
header(s, "我们的方案", "一句话:让 AI 帮你找对标")
text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.8),
     [[("上传报关单 → AI 自动推荐", 26, INK, True), ("「可用」的亚马逊对标链接", 26, ACCENT, True)],
      [("并对每一条说清为什么是它;用得越多,知识库越准、越快。", 18, MUTED, False)]], space=14)
for i, (h, b) in enumerate([
    ("不用 RPA", "用 Keepa 官方数据,不碰反爬"),
    ("有据可依", "价格 / 销量 / 评论 + AI 判断"),
    ("越用越准", "用户反馈沉淀长期知识库"),
]):
    card(s, Inches(0.9 + i * 3.95), Inches(4.4), Inches(3.7), Inches(1.7), h, b)
footer(s, 3)

# ───────── S4 工作流程 ─────────
s = slide()
header(s, "怎么工作", "全流程自动,人只做最后确认")
steps = ["上传表格", "AI 解析归一", "AI 翻译关键词", "Keepa 检索", "AI 相似度分析",
         "三重核验", "按要求择优", "推荐可用链接", "用户反馈入库", "反哺推荐"]
x0, y0, cw, gap = Inches(0.8), Inches(2.4), Inches(2.15), Inches(0.18)
per = 4
for i, st in enumerate(steps):
    row, col = divmod(i, per)
    l = x0 + col * (cw + gap)
    t = y0 + row * Inches(1.15)
    fill = ACCENT if st in ("推荐可用链接",) else ACCENTSOFT
    fg = WHITE if st in ("推荐可用链接",) else ACCENT
    chip(s, l, t, cw, st, fill=fill, fg=fg)
    if col < per - 1 and i < len(steps) - 1:
        text(s, l + cw, t, gap, Inches(0.62), [[("→", 16, FAINT, True)]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.7),
     [[("亮点:三重核验(商标/材质/原产地)+ 按你的要求择优(最便宜/目标价)+ 反馈入库,越用越准。", 14, MUTED, False)]])
footer(s, 4)

# ───────── S5 三大核心能力 ─────────
s = slide()
header(s, "核心能力", "三件事,做到极致")
cards = [
    ("AI 解析表格", "报关单 / 装箱单(Excel / PDF / 图片)自动解析、字段归一;中英混排、只给材质也能读懂。", "1"),
    ("AI 推荐可用链接", "电商翻译 → Keepa 检索 → 相似度分析 → 三重核验 → 按要求择优,推荐含价格/销量/评论的可用对标。", "2"),
    ("长期知识库", "用户采纳/否决持续回流入库,沉淀客户专属对标知识库,同类下次秒级命中,越用越准越快。", "3"),
]
for i, (h, b, n) in enumerate(cards):
    card(s, Inches(0.8 + i * 4.0), Inches(2.3), Inches(3.75), Inches(3.4), h, b, num=n)
footer(s, 5)

# ───────── S6 为什么更聪明 ─────────
s = slide()
header(s, "差异化", "比几千块的普通脚本,聪明在哪")
rows = [
    ("翻译", "词典逐词直译,模板一变就崩", "电商语境翻译,买家真实搜索词"),
    ("查找对标", "关键词精确匹配,差一个字就找不到", "语义检索 + Keepa 数据,相似即命中"),
    ("判断同款", "撞词就算 / 人工肉眼一个个看", "多模态相关性,自动排掉不同款"),
    ("商标合规", "多数不查,漏检 = 合规风险", "AI 揪品牌词 → EUIPO 商标库核验"),
    ("适应性", "页面/模板一改就失效,常维护", "自适应,稳定、维护轻"),
    ("可解释", "黑箱,出错要人收拾", "每条对标都给详细理由"),
]
y = Inches(2.15)
text(s, Inches(4.2), Inches(1.85), Inches(4), Inches(0.3), [[("普通脚本", 12, FAINT, True)]])
text(s, Inches(8.6), Inches(1.85), Inches(4), Inches(0.3), [[("我们的 AI", 12, ACCENT, True)]])
for dim, a, b in rows:
    text(s, Inches(0.9), y, Inches(3.0), Inches(0.6), [[(dim, 14, INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.2), y, Inches(4.2), Inches(0.6), [[("✕ ", 12, DANGER, True), (a, 12.5, MUTED, False)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(8.6), y, Inches(4.2), Inches(0.6), [[("✓ ", 12, OK, True), (b, 12.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.72)
footer(s, 6)

# ───────── S7 不用 RPA ─────────
s = slide()
header(s, "技术选型", "不用 RPA:Keepa API + AI")
bullets(s, [
    ("合法、稳定、可预算", "用 Keepa 官方商用 API,不碰反爬、不踩验证码、无封号风险"),
    ("上线快、维护轻", "无需住宅代理与反爬对抗,合规清晰"),
    ("数据天生适配", "Keepa 自带价格历史 + 销量排名(BSR)+ 评论,正好满足销量/评分/最便宜的选品"),
], t=Inches(2.1), gap=12)
rect(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(1.0), fill=ACCENTSOFT, radius=True)
text(s, Inches(1.2), Inches(5.5), Inches(11), Inches(1.0),
     [[("成本对照:自建 RPA 首年约 ¥22 万 – ¥54 万 ", 16, MUTED, False),
       ("→ 本方案首年 ≈ ¥3.1 万", 16, ACCENT, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)

# ───────── S8 自定义要求 ─────────
s = slide()
header(s, "可自定义", "只推符合你要求的,再选最划算的")
for i, (h, b) in enumerate([
    ("硬性要求", "原产地中国 · 排除品牌 · 材质匹配"),
    ("条件门槛", "最低评论数(销量)· 最低评分"),
    ("选品目标", "最便宜 或 最接近目标价"),
]):
    card(s, Inches(0.9 + i * 3.95), Inches(2.3), Inches(3.7), Inches(1.9), h, b)
text(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(1.5),
     [[("先按要求筛掉不合格候选,再在合格里按目标择优。", 17, INK, True)],
      [("要求可一键存为预设(稳妥合规 / 成本优先 / 对标某价位…);也支持自然语言:「找最便宜的」「目标价 €15」「评论过千」。", 14, MUTED, False)]], space=12)
footer(s, 8)

# ───────── S9 知识库资产 ─────────
s = slide()
header(s, "数据资产", "越用越值钱:你的专属对标知识库")
bullets(s, [
    ("反馈闭环", "每次采纳 / 否决都回流入库,持续校准 AI 推荐"),
    ("越用越快越准", "同类商品下次秒级命中,命中率随使用上升"),
    ("沉淀为资产", "形成客户专属对标库,数据越厚壁垒越高、粘性越强"),
], t=Inches(2.1), gap=14)
footer(s, 9)

# ───────── S10 投资回报 ─────────
s = slide()
header(s, "投资回报", "省时、省人、零漏检")
stats = [("~20 分钟", "人工 / 件"), ("秒级", "系统 / 件"), ("≈ 80×", "效率提升"), ("0", "核验漏检")]
for i, (big, lab) in enumerate(stats):
    l = Inches(0.85 + i * 3.0)
    rect(s, l, Inches(2.3), Inches(2.8), Inches(2.0), fill=WHITE, line=LINE, radius=True)
    text(s, l, Inches(2.5), Inches(2.8), Inches(1.0), [[(big, 34, ACCENT, True)]], align=PP_ALIGN.CENTER)
    text(s, l, Inches(3.6), Inches(2.8), Inches(0.5), [[(lab, 14, MUTED, False)]], align=PP_ALIGN.CENTER)
rect(s, Inches(0.85), Inches(4.7), Inches(11.5), Inches(1.0), fill=ACCENTSOFT, radius=True)
text(s, Inches(1.2), Inches(4.7), Inches(11), Inches(1.0),
     [[("商标、材质、原产地一个都不会漏查;全天候不间断 · 首年总投入 ≈ ¥3.1 万。", 16, ACCENT, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, 10)

# ───────── S11 报价与交付 ─────────
s = slide()
header(s, "报价与交付", "一口价,零门槛先跑起来")
bullets(s, [
    ("一次性开发(打包,核心 7 项):¥15,000", "表格解析 / 对标推荐 / 三重核验 / 选品引擎 / 数据库 / 知识库 / 部署培训"),
    ("月度运营:约 ¥1,400 – ¥4,300", "Keepa 订阅 + AI 调用 + 服务器(商标核验用 EUIPO 免费 API)"),
    ("🎁 赠送前 3 个月运营 · 交付 4–6 周 · 付款 50/30/20", "首年总投入 ≈ ¥3.1 万"),
], t=Inches(2.1), gap=14)
footer(s, 11)

# ───────── S12 可扩展 ─────────
s = slide()
header(s, "路线图", "可分期扩展(架构预留,核心不重构)")
exts = [
    ("AI 自动填表 + 审核台", "批量填表、逐条通过/驳回、自然语言操作"),
    ("工单接入", "飞书 / 企业微信 / 邮箱自动收单、读聊天与附件"),
    ("结果导出 / 回写", "CSV 导出、回写客户数据库 / ERP"),
    ("多站点扩展", "亚马逊其他国家站点"),
]
for i, (h, b) in enumerate(exts):
    row, col = divmod(i, 2)
    card(s, Inches(0.9 + col * 5.9), Inches(2.3 + row * 1.9), Inches(5.6), Inches(1.6), h, b)
footer(s, 12)

# ───────── S13 CTA ─────────
s = slide(ACCENT)
rect(s, 0, 0, SW, SH, fill=ACCENT)
text(s, Inches(1.0), Inches(2.5), Inches(11.4), Inches(1.2),
     [[("现场看效果,90 秒看懂价值", 36, WHITE, True)]])
text(s, Inches(1.0), Inches(3.9), Inches(11.4), Inches(0.6),
     [[("在线演示:https://amazon-linkfinder-demo.vercel.app", 18, RGBColor(0xE7, 0xE9, 0xFB), True)]])
text(s, Inches(1.0), Inches(5.2), Inches(11.4), Inches(0.6),
     [[("联系:[联系人] · [电话] · [邮箱]", 15, RGBColor(0xC7, 0xCB, 0xF2), False)]])

out = "销售PPT-Amazon对标选品系统.pptx"
prs.save(out)
print("PPT 生成:", out, "·", len(prs.slides.__iter__.__self__._sldIdLst), "页")
